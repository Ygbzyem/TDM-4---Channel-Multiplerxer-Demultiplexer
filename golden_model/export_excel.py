"""Write the verification results to a professional Excel workbook.

    verification_results.xlsx   (every run is APPENDED, nothing is overwritten)
      Summary  - latest run (totals, latency in cycles and ns, chart) and the cumulative totals of
                 all runs; live formulas
      Results  - one row per input vector of EVERY run (Run #, date/time, input file, inputs,
                 golden vs RTL outputs and bit streams, latency, PASS/FAIL); every check is an
                 Excel formula; filter by "Run #" to look at a single run
      History  - one row per run (date/time, input, totals, latency)

Needs the third-party package openpyxl:   python -m pip install openpyxl
It is optional: run_verification.py prints a hint and carries on when it is missing.
"""

from datetime import datetime
from pathlib import Path

CLOCK_PERIOD_NS = 10            # tb_tdm_top_filebased.v: always #5 clk = ~clk  -> 10 ns
FONT = "Arial"

# fixed cells on the Summary sheet that the Results formulas refer to
CLOCK_CELL = "Summary!$B$11"
MAXLAT_CELL = "Summary!$B$12"

RESULT_HEADS = [
    ("Run #", 7), ("Run date / time", 18), ("Input file", 22), ("Vector #", 9), ("Type", 12),
    ("Input (hex)", 13), ("CH0 in", 8), ("CH1 in", 8), ("CH2 in", 8), ("CH3 in", 8),
    ("Golden CHx_OUT (hex)", 16), ("RTL CHx_OUT (hex)", 16), ("Output check", 12),
    ("Golden bit stream (32 bit)", 36), ("RTL bit stream (32 bit)", 36), ("Stream check", 12),
    ("Latency (cycles)", 11), ("Latency (ns)", 11), ("Latency check", 13), ("Result", 10),
]
HISTORY_HEADS = [
    ("Run #", 7), ("Run date / time", 20), ("Input file", 24), ("Simulator / tool", 28),
    ("Vectors", 9), ("Passed", 9), ("Failed", 9), ("Cold start latency (cycles)", 16),
    ("Max latency (cycles)", 14), ("Overall", 10),
]


def _styles():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin = Side(style="thin", color="B8C2CC")
    return {
        "title": Font(name=FONT, size=15, bold=True, color="1F3A5F"),
        "section": Font(name=FONT, size=11, bold=True, color="FFFFFF"),
        "section_fill": PatternFill("solid", fgColor="2E5E8C"),
        "head": Font(name=FONT, size=10, bold=True, color="FFFFFF"),
        "head_fill": PatternFill("solid", fgColor="1F3A5F"),
        "body": Font(name=FONT, size=10),
        "bold": Font(name=FONT, size=10, bold=True),
        "mono": Font(name="Consolas", size=10),
        "input": Font(name=FONT, size=10, color="0000FF"),
        "input_fill": PatternFill("solid", fgColor="FFFF99"),
        "band": PatternFill("solid", fgColor="EEF3F8"),
        "note": Font(name=FONT, size=9, italic=True, color="666666"),
        "border": Border(left=thin, right=thin, top=thin, bottom=thin),
        "center": Alignment(horizontal="center", vertical="center"),
        "left": Alignment(horizontal="left", vertical="center"),
        "wrap": Alignment(horizontal="center", vertical="center", wrap_text=True),
    }


def _add_status_formatting(ws, cell_range):
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Font, PatternFill

    good = dict(font=Font(name=FONT, size=10, bold=True, color="006100"),
                fill=PatternFill("solid", bgColor="C6EFCE", fgColor="C6EFCE"))
    bad = dict(font=Font(name=FONT, size=10, bold=True, color="9C0006"),
               fill=PatternFill("solid", bgColor="FFC7CE", fgColor="FFC7CE"))
    for word in ('"PASS"', '"MATCH"', '"OK"'):
        ws.conditional_formatting.add(cell_range, CellIsRule(operator="equal", formula=[word], **good))
    for word in ('"FAIL"', '"MISMATCH"', '"OVER LIMIT"', '"NEVER"'):
        ws.conditional_formatting.add(cell_range, CellIsRule(operator="equal", formula=[word], **bad))


def _write_header(ws, st, heads):
    for c, (h, w) in enumerate(heads, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font, cell.fill, cell.alignment, cell.border = st["head"], st["head_fill"], st["wrap"], st["border"]
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[1].height = 32


def _append_results(ws, st, results, run_no, stamp, input_file):
    """Append one block of rows (this run) below the existing ones. Returns (first_row, last_row)."""
    first = ws.max_row + 1
    for i, r in enumerate(results):
        row = first + i
        vals = {
            "A": run_no, "B": stamp, "C": input_file,
            "D": r["index"],
            "E": "Cold start" if r["index"] == 0 else "Mid-stream",
            "F": r["input"].upper(),
            "G": f"=MID(F{row},1,2)", "H": f"=MID(F{row},3,2)",
            "I": f"=MID(F{row},5,2)", "J": f"=MID(F{row},7,2)",
            "K": r["golden_out"].upper(),
            "L": r["rtl_out"].upper(),
            "M": f'=IF(K{row}=L{row},"MATCH","MISMATCH")',
            "N": r["golden_stream"],
            "O": r["rtl_stream"],
            "P": f'=IF(N{row}=O{row},"MATCH","MISMATCH")',
            "Q": r["latency"],
            "R": f'=IF(Q{row}<0,"n/a",Q{row}*{CLOCK_CELL})',
            "S": f'=IF(Q{row}<0,"NEVER",IF(Q{row}<={MAXLAT_CELL},"OK","OVER LIMIT"))',
            "T": f'=IF(AND(M{row}="MATCH",P{row}="MATCH",S{row}="OK"),"PASS","FAIL")',
        }
        for col, v in vals.items():
            cell = ws[f"{col}{row}"]
            cell.value = v
            cell.border = st["border"]
            cell.font = st["mono"] if col in "FGHIJKLNO" else st["body"]
            cell.alignment = st["left"] if col in "BCNO" else st["center"]
            if run_no % 2 == 0:                       # light band so runs are easy to tell apart
                cell.fill = st["band"]
        ws[f"R{row}"].number_format = "0"
    return first, ws.max_row


def _build_summary(ws, st, summary, meta, run_no, first, last, clock, maxlat):
    from openpyxl.chart import BarChart, Reference
    from openpyxl.comments import Comment

    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 62
    ws.column_dimensions["C"].width = 16
    ws["A1"] = "TDM 4-Channel Mux/Demux - RTL vs Golden Model Verification"
    ws["A1"].font = st["title"]

    def section(row, text):
        for col in "ABC":
            ws[f"{col}{row}"].fill = st["section_fill"]
        ws[f"A{row}"] = text
        ws[f"A{row}"].font = st["section"]

    def line(row, label, value, fmt=None, bold=False):
        ws[f"A{row}"] = label
        ws[f"A{row}"].font = st["bold"]
        ws[f"B{row}"] = value
        ws[f"B{row}"].font = st["bold"] if bold else st["body"]
        ws[f"B{row}"].alignment = st["left"]
        for col in "AB":
            ws[f"{col}{row}"].border = st["border"]
        if fmt:
            ws[f"B{row}"].number_format = fmt

    allr = lambda col: f"Results!${col}$2:${col}${last}"          # every run
    cur = lambda col: f"Results!${col}${first}:${col}${last}"     # latest run only

    section(3, "Run information (latest run)")
    line(4, "Latest run #", run_no, "0")
    line(5, "Run date / time", meta["timestamp"])
    line(6, "Input file", summary["input_file"])
    line(7, "Simulator / tool", meta.get("simulator", ""))
    line(8, "Reference model", "Python golden model (golden_model.TDMGoldenModel)")
    line(9, "Design under test", "tdm_top (tdm_counter, tdm_mux, tx_shift_reg, rx_shift_reg, tdm_demux)")
    line(10, "Runs recorded in this workbook", f"=MAX({allr('A')})", "0")
    line(11, "Clock period (ns)", clock)
    line(12, "Max allowed latency (cycles)", maxlat)
    for r in (11, 12):
        ws[f"B{r}"].font, ws[f"B{r}"].fill = st["input"], st["input_fill"]
        ws[f"C{r}"] = "<- editable"
        ws[f"C{r}"].font = st["note"]
    ws["B11"].comment = Comment("Assumption: tb_tdm_top_filebased.v generates the clock with "
                                "'always #5 clk = ~clk' and timescale 1ns/1ps, i.e. 10 ns. "
                                "Edit this cell if the testbench clock changes.", "golden_model")
    ws["B12"].comment = Comment("Source: golden_model/latency_checker.py MAX_LATENCY = 41 "
                                "(31 wait for load slot + 8 serialise + 2 output registers; "
                                "verified by sweeping all 32 bit_count phases).", "golden_model")

    section(14, "Latest run - verification summary")
    line(15, "Total vectors", f"=COUNTA({cur('D')})", "0")
    line(16, "Vectors passed", f'=COUNTIF({cur("T")},"PASS")', "0")
    line(17, "Vectors failed", "=B15-B16", "0")
    line(18, "Pass rate", "=IF(B15=0,0,B16/B15)", "0.0%")
    line(19, "Overall status", '=IF(AND(B15>0,B17=0,B28="PASS"),"PASS","FAIL")', bold=True)

    section(21, "Latest run - latency")
    ws["B21"], ws["C21"] = "Cycles", "ns"
    for c in ("B21", "C21"):
        ws[c].font, ws[c].alignment = st["section"], st["center"]
    has_mid = last - first >= 1
    first_mid = first + 1
    rows = [
        (22, "Cold start (vector 0)", f"=Results!Q{first}"),
        (23, "Mid-stream minimum", f"=MIN(Results!$Q${first_mid}:$Q${last})" if has_mid else "n/a"),
        (24, "Mid-stream maximum", f"=MAX(Results!$Q${first_mid}:$Q${last})" if has_mid else "n/a"),
        (25, "Mid-stream average", f"=AVERAGE(Results!$Q${first_mid}:$Q${last})" if has_mid else "n/a"),
        (26, "Worst case of all vectors", f"=MAX({cur('Q')})"),
        (27, "Allowed maximum", "=B12"),
        (28, "Latency check", f'=IF(AND(MIN({cur("Q")})>=0,MAX({cur("Q")})<=B12),"PASS","FAIL")'),
    ]
    for r, label, f in rows:
        ws[f"A{r}"], ws[f"B{r}"] = label, f
        ws[f"A{r}"].font = st["bold"]
        ws[f"B{r}"].font = st["bold"] if r == 28 else st["body"]
        ws[f"B{r}"].alignment = st["center"]
        ws[f"A{r}"].border = ws[f"B{r}"].border = st["border"]
        if r < 28:
            ws[f"B{r}"].number_format = "0.0" if r == 25 else "0"
            ws[f"C{r}"] = f'=IF(ISNUMBER(B{r}),B{r}*$B$11,"n/a")'
            ws[f"C{r}"].font, ws[f"C{r}"].alignment, ws[f"C{r}"].border = st["body"], st["center"], st["border"]
            ws[f"C{r}"].number_format = "0.0"

    section(30, "All runs (cumulative)")
    line(31, "Runs recorded", "=B10", "0")
    line(32, "Total vectors checked", f"=COUNTA({allr('D')})", "0")
    line(33, "Vectors passed", f'=COUNTIF({allr("T")},"PASS")', "0")
    line(34, "Vectors failed", "=B32-B33", "0")
    line(35, "Pass rate", "=IF(B32=0,0,B33/B32)", "0.0%")
    line(36, "Worst latency of all runs (cycles)", f"=MAX({allr('Q')})", "0")
    line(37, "Average mid-stream latency (cycles)",
         f'=IFERROR(AVERAGEIF({allr("E")},"Mid-stream",{allr("Q")}),"n/a")', "0.0")
    line(38, "Overall status (all runs)",
         f'=IF(AND(B32>0,B34=0,MIN({allr("Q")})>=0,B36<=B12),"PASS","FAIL")', bold=True)
    for r in (19, 28, 38):
        ws[f"B{r}"].alignment = st["center"] if r == 28 else st["left"]
    _add_status_formatting(ws, "B19")
    _add_status_formatting(ws, "B28")
    _add_status_formatting(ws, "B38")

    ws["A40"] = ("Every run is appended to the Results sheet (filter by Run #); the sections above are formulas "
                 "over the Results sheet. Vector 0 of a run is a cold start (latency counted from reset release); "
                 "later vectors change the inputs mid-stream at bit_count = 24. If values look empty, click "
                 "'Enable Editing' in Excel.")
    ws["A40"].font = st["note"]

    ch = BarChart()
    ch.type = "col"
    ch.title = f"Latency per input vector - run #{run_no}"
    ch.y_axis.title = "Latency (clock cycles)"
    ch.x_axis.title = "Vector #"
    res_ws = summary["_results_ws"]
    ch.add_data(Reference(res_ws, min_col=17, min_row=first, max_row=last), titles_from_data=False)
    ch.set_categories(Reference(res_ws, min_col=4, min_row=first, max_row=last))
    ch.legend = None
    ch.height, ch.width = 8.5, 17
    ch.y_axis.delete = False
    ch.x_axis.delete = False
    ws.add_chart(ch, "E3")


def _append_history(ws, st, summary, meta, run_no):
    from openpyxl.formatting.formatting import ConditionalFormattingList
    results, lat = summary["results"], summary["latency"]
    passed = sum(1 for r in results if r["ok"])
    row = ws.max_row + 1
    values = [run_no, meta["timestamp"], summary["input_file"], meta.get("simulator", ""),
              len(results), passed, len(results) - passed, lat["cold_start"],
              max((r["latency"] for r in results), default=None),
              "PASS" if summary["all_pass"] else "FAIL"]
    for c, v in enumerate(values, start=1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.font, cell.border = st["body"], st["border"]
        cell.alignment = st["left"] if c in (2, 3, 4) else st["center"]
    ws.auto_filter.ref = f"A1:J{row}"
    ws.conditional_formatting = ConditionalFormattingList()
    _add_status_formatting(ws, f"J2:J{max(row, 2)}")


def write_excel(path, summary, simulator=""):
    """Append this run to the workbook (created on first use). summary = same dict as
    report_generator.write_report (input_file, results, latency, all_pass). Returns the path."""

    from openpyxl import Workbook, load_workbook
    from openpyxl.formatting.formatting import ConditionalFormattingList
    from openpyxl.worksheet.properties import PageSetupProperties

    path = Path(path)
    st = _styles()
    meta = {"timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"), "simulator": simulator}
    clock, maxlat = CLOCK_PERIOD_NS, summary["latency"]["max_allowed"]

    if path.exists():
        wb = load_workbook(path)
        if "Summary" in wb.sheetnames:                  # keep the user's edits of the yellow cells
            old = wb["Summary"]
            for cell, name in (("B11", "clock"), ("B12", "maxlat")):
                if isinstance(old[cell].value, (int, float)):
                    if name == "clock":
                        clock = old[cell].value
                    else:
                        maxlat = old[cell].value
            del wb["Summary"]
        for name, first_head in (("Results", "Run #"), ("History", "Run #")):
            if name in wb.sheetnames and wb[name]["A1"].value != first_head:
                del wb[name]                              # sheet from an older layout: start it again
    else:
        wb = Workbook()
        wb.remove(wb.active)

    ws_sum = wb.create_sheet("Summary", 0)
    if "Results" in wb.sheetnames:
        ws_res = wb["Results"]
    else:
        ws_res = wb.create_sheet("Results", 1)
        _write_header(ws_res, st, RESULT_HEADS)
        ws_res.freeze_panes = "F2"
    if "History" in wb.sheetnames:
        ws_his = wb["History"]
    else:
        ws_his = wb.create_sheet("History")
        _write_header(ws_his, st, HISTORY_HEADS)
        ws_his.freeze_panes = "A2"

    previous = [c.value for c in ws_res["A"][1:] if isinstance(c.value, (int, float))]
    run_no = int(max(previous)) + 1 if previous else 1

    first, last = _append_results(ws_res, st, summary["results"], run_no, meta["timestamp"], summary["input_file"])
    ws_res.auto_filter.ref = f"A1:T{last}"
    ws_res.conditional_formatting = ConditionalFormattingList()      # no duplicate rules after many runs
    for rng in (f"M2:M{last}", f"P2:P{last}", f"S2:T{last}"):
        _add_status_formatting(ws_res, rng)

    summary = dict(summary, _results_ws=ws_res)
    _build_summary(ws_sum, st, summary, meta, run_no, first, last, clock, maxlat)
    _append_history(ws_his, st, summary, meta, run_no)

    for ws in (ws_sum, ws_res, ws_his):                     # print: landscape, one page wide
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws_res.print_title_rows = "1:1"
    ws_sum.sheet_properties.tabColor = "1F3A5F"
    ws_res.sheet_properties.tabColor = "2E5E8C"
    wb.active = 0
    wb.save(path)
    return path
